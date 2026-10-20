#!/usr/bin/env python3
"""Classroom deck for the film 圈子从哪里来, built from the same engine data as the film.

One slide description drives three outputs:
  site/downloads/label-propagation.pptx  editable (native text, shapes, tables, speaker notes)
  site/downloads/label-propagation.pdf   vector preview (Pycairo, embedded Noto CJK)
  site/slides/slide-NN.png               web preview, plus presentation/notes.json and site/downloads/notes.json

  python3 scripts/build_deck.py
"""
import json, math, subprocess
from pathlib import Path
import cairo
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
W, H = 1280, 720                     # slide units; 1 unit = 9525 EMU (one pixel at 96 dpi)
EMU = 9525
PNG_SCALE = 1.5                      # 1920 x 1080 previews

# Film palette (scripts/render_film_v3.py)
BG, PANEL, LINE = '#061219', '#0d222c', '#24414b'
IVORY, MUTED, FAINT = '#eeeadd', '#8fa3a3', '#3b5560'
GOLD, TEAL, BLUE, ROSE, VIOLET = '#ddbd78', '#76cbb5', '#86a9e0', '#d99a9a', '#b9a0dc'
LABEL = {'A': GOLD, 'B': TEAL, 'C': ROSE, 'D': BLUE, 'E': VIOLET}
GROUP = [TEAL, GOLD, BLUE, ROSE, VIOLET]
SANS_PNG, SERIF_PNG = 'Noto Sans CJK SC', 'Noto Serif CJK SC'
SANS_PPT, SERIF_PPT = 'PingFang SC', 'Songti SC'

TEACHER = 'https://github.com/yongfu-li/eda-learning/tree/main/courses/learn_clustering/module02-01-label-propagation'
PAPER = 'https://arxiv.org/abs/0709.2938'
SITE = 'https://su3037.github.io/clustering-lab5-label-propagation/'

FILM = json.loads((ROOT / 'data/film-v3.json').read_text())
SITE_GRAPHS = json.loads(subprocess.check_output(
    ['node', '--input-type=module', '-e',
     "const m=await import('./site/film-graphs.js');console.log(JSON.stringify({contacts:m.CONTACTS,chip:m.CHIP}))"],
    cwd=ROOT, text=True))
CONTACTS, CHIP = SITE_GRAPHS['contacts'], SITE_GRAPHS['chip']
GOLDEN = FILM['golden']
TEACHER_EDGES = [('A', 'B', 5), ('A', 'C', 1), ('B', 'C', 4), ('C', 'D', 2), ('D', 'E', 5), ('C', 'E', 1)]
TEACHER_POS = {'A': (150, 250), 'B': (390, 200), 'C': (310, 405), 'D': (620, 420), 'E': (505, 585)}


def rgb(c): c = c.lstrip('#'); return tuple(int(c[i:i + 2], 16) / 255 for i in (0, 2, 4))
def mix(c, a, under=BG):
    """Pre-blend a colour over the background so both outputs need no transparency."""
    x, y = rgb(c), rgb(under)
    return '#' + ''.join(f'{round((p * a + q * (1 - a)) * 255):02x}' for p, q in zip(x, y))


# ---------------------------------------------------------------- slide model
class Slide:
    def __init__(self, title, notes, chapter=''):
        self.title, self.notes, self.chapter, self.items = title, notes, chapter, []
    def add(self, kind, **kw): self.items.append((kind, kw)); return self
    def text(self, x, y, w, h, s, size=24, color=IVORY, bold=False, serif=False, align='l', valign='t', lh=1.35):
        return self.add('text', x=x, y=y, w=w, h=h, s=s, size=size, color=color, bold=bold, serif=serif, align=align, valign=valign, lh=lh)
    def line(self, x1, y1, x2, y2, color=LINE, width=2): return self.add('line', x1=x1, y1=y1, x2=x2, y2=y2, color=color, width=width)
    def circle(self, cx, cy, r, fill=None, stroke=None, width=2): return self.add('circle', cx=cx, cy=cy, r=r, fill=fill, stroke=stroke, width=width)
    def rect(self, x, y, w, h, fill=None, stroke=None, width=1.5, radius=10): return self.add('rect', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, width=width, radius=radius)
    def table(self, x, y, widths, rowh, rows, size=21, accent=None):
        return self.add('table', x=x, y=y, widths=widths, rowh=rowh, rows=rows, size=size, accent=accent or {})
    def image(self, path, x, y, w, h): return self.add('image', path=path, x=x, y=y, w=w, h=h)


SLIDES = []
def slide(title, notes, chapter=''):
    s = Slide(title, notes, chapter); SLIDES.append(s)
    if title:
        s.text(64, 40, 1000, 64, title, 38, IVORY, serif=True)
        s.line(64, 112, 184, 112, GOLD, 2)
    return s


def frame(s, i):
    if s.title:
        if s.chapter: s.text(980, 52, 236, 30, s.chapter, 16, MUTED, align='r')
        s.text(64, 680, 700, 24, 'Su3037 · Clustering Lab 05 · 圈子从哪里来', 14, FAINT)
        s.text(1150, 680, 66, 24, f'{i:02d}', 14, FAINT, align='r')


# ---------------------------------------------------------------- graph helpers
def fit(pos, box):
    xs = [p[0] for p in pos.values()]; ys = [p[1] for p in pos.values()]
    x0, y0, x1, y1 = box
    k = min((x1 - x0) / (max(xs) - min(xs) or 1), (y1 - y0) / (max(ys) - min(ys) or 1))
    ox = x0 + ((x1 - x0) - (max(xs) - min(xs)) * k) / 2; oy = y0 + ((y1 - y0) - (max(ys) - min(ys)) * k) / 2
    return {n: (ox + (p[0] - min(xs)) * k, oy + (p[1] - min(ys)) * k) for n, p in pos.items()}


def ranked_colours(labels, nodes):
    """Colour groups by size, then by first member, so the biggest group is always teal."""
    groups = {}
    for n in nodes: groups.setdefault(labels[n], []).append(n)
    order = sorted(groups, key=lambda L: (-len(groups[L]), nodes.index(groups[L][0])))
    return {L: GROUP[i % len(GROUP)] for i, L in enumerate(order)}


def big_graph(s, g, pos, labels=None, names=None, cut=False, r=8, weight_scale=.45):
    col = ranked_colours(labels, g['nodes']) if labels else {}
    for u, v, w in g['edges']:
        a, b = pos[u], pos[v]
        crossing = labels and labels[u] != labels[v]
        if cut and crossing: s.line(*a, *b, GOLD, 1.2 + w * .5)
        elif labels and not crossing: s.line(*a, *b, mix(col[labels[u]], .45), .8 + w * weight_scale)
        else: s.line(*a, *b, mix(MUTED, .28 if labels else .4), .8 + w * weight_scale)
    for n in g['nodes']:
        x, y = pos[n]; c = col[labels[n]] if labels else MUTED
        s.circle(x, y, r, BG, c, 2.4)
        if labels: s.circle(x, y, r * .5, c)
        if names: s.text(x - 40, y - r - 21, 80, 20, names[n], 13, c if labels else MUTED, align='c')


def teacher_graph(s, labels, cut=False, edges=TEACHER_EDGES, pos=TEACHER_POS, highlight=(), dx=0, dy=0):
    pos = {n: (x + dx, y + dy) for n, (x, y) in pos.items()}
    for u, v, w in edges:
        a, b = pos[u], pos[v]
        crossing = labels[u] != labels[v]
        c = GOLD if cut and crossing else (TEAL if (u, v) in highlight or (v, u) in highlight else mix(MUTED, .55))
        s.line(*a, *b, c, 1.5 + w * .7)
        mx, my = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
        s.circle(mx, my, 15, BG, c, 1.2)
        s.text(mx - 15, my - 14, 30, 26, str(w), 17, c, bold=True, align='c', valign='m')
    for n, (x, y) in pos.items():
        c = LABEL[labels[n]]
        s.circle(x, y, 32, BG, c, 3.5)
        s.text(x - 32, y - 22, 64, 40, n, 28, IVORY, bold=True, align='c', valign='m')
        s.rect(x - 44, y + 40, 88, 28, mix(c, .18), None, radius=14)
        s.text(x - 44, y + 41, 88, 26, '标签 ' + labels[n], 16, c, align='c', valign='m')


def stat(s, x, y, w, big, small, color=GOLD):
    s.text(x, y, w, 56, big, 40, color, bold=True, serif=True)
    s.text(x, y + 58, w, 30, small, 17, MUTED)


def bullets(s, x, y, w, items, size=24, gap=1.0, color=IVORY, mark=GOLD):
    for it in items:
        lines = sum(max(1, math.ceil(len(p) * size * .95 / (w - 24))) for p in it.split('\n'))
        s.circle(x + 6, y + size * .68, 4, mark)
        s.text(x + 24, y, w - 24, size * 1.35 * lines, it, size, color)
        y += size * 1.35 * lines + size * .9 * gap


# ---------------------------------------------------------------- the deck
first_round = GOLDEN['events'][:5]
c_event = GOLDEN['events'][2]
lp_chip, ruler_chip = FILM['lp'], FILM['geometric']
social = FILM['social']['run']
assert social['rounds'] == 3 and social['metrics']['cut'] == 4 and sorted(social['metrics']['sizes']) == [8, 10, 10]
assert lp_chip['rounds'] == 5 and lp_chip['metrics']['cut'] == 2 and ruler_chip['metrics']['cut'] == 88 and FILM['initialCut'] == 202
assert GOLDEN['rounds'] == 2 and GOLDEN['metrics']['cut'] == 3 and FILM['reverse']['rounds'] == 3
assert FILM['strongCD']['metrics']['cut'] == 5 and FILM['pairSync']['status'] == 'cycle'

# 01 cover
s = slide('', '开场播放短片《圈子从哪里来》：6 分 20 秒，1080p，无配音，中英双语字幕。上篇（00:00–03:09）讲为什么需要标签传播，下篇（03:09–05:43）分九层讲它怎样工作，最后讲现实中的用法。课堂时间紧时可分两次播放：先放上篇，讲完第 7–15 页再放下篇对照。\n\n配乐为用户指定的《We Never (Remix)（若如初见）》Hi Noise / 秋成，不属于本项目的开放许可。通讯录与 24 单元电路都是虚构的教学图。作者 Su3037。')
s.image('site/media/circles-cover.jpg', 0, 0, W, H)

# 02 contacts
s = slide('通讯录里的圈子', '短片 00:00。先问：打开手机通讯录，几百个名字，怎样分成几个圈子？我们不需要知道每个人是谁，只看谁和谁常联系。好的分组是圈内联系紧、圈与圈之间联系少。网络科学把这件事叫社群发现（community detection）。这里的 28 人和 61 条联系是虚构的教学数据，只用来解释图结构，不用来判断真实的人。', '短片 00:00')
big_graph(s, CONTACTS, fit(CONTACTS['positions'], (70, 150, 730, 640)), names=CONTACTS['names'])
s.text(800, 170, 420, 50, '几百个名字，', 30, IVORY, serif=True)
s.text(800, 215, 420, 50, '怎么分成几个圈子？', 30, IVORY, serif=True)
bullets(s, 800, 310, 420, ['只看谁和谁常联系', '好的分组：圈内紧，圈间少', '这件事叫社群发现\ncommunity detection'], 22)

# 03 the 2007 idea on contacts
s = slide('一个极简的办法', '短片 00:34。Raghavan、Albert、Kumara 在 2007 年提出标签传播：每个人看看朋友们现在属于哪个圈子，加入联系最多的那一个，一遍遍重复，圈子自己浮现。右边的数字是项目算法在这张虚构通讯录上实际运行的结果（异步、最小标签平票）：3 轮停下，分成 8、10、10 三个圈子，跨圈只切开权重 4（总权重 238）。', '短片 00:34')
big_graph(s, CONTACTS, fit(CONTACTS['gathered'], (70, 150, 730, 640)), social['labels'], CONTACTS['names'])
s.text(800, 160, 420, 90, '每个人看看朋友们现在属于哪个圈子，\n加入联系最多的那一个。', 23, IVORY, lh=1.5)
s.text(800, 262, 420, 30, 'Raghavan · Albert · Kumara，2007', 17, MUTED)
stat(s, 800, 330, 200, '3 轮', '停下，没有人再改变')
stat(s, 1010, 330, 210, '8+10+10', '浮现出三个圈子')
stat(s, 800, 470, 420, '只切开 4', '跨圈的联系，权重合计（共 238）', TEAL)

# 04 the chip
s = slide('芯片里的同一个问题', '短片 00:51。芯片有几百万个单元，由导线连成一张网。设计时，EDA 软件要先把它们分成几块：块内用短导线，跨块只能靠长导线。导线越长，信号越慢、越耗电，也占布线资源。所以目标和通讯录一样：块内紧，块间少。左图是 24 个单元、56 条连线的虚构电路，线越粗代表连接越强。', '短片 00:51')
big_graph(s, CHIP, fit(CHIP['positions'], (70, 150, 730, 640)), r=10, weight_scale=.6)
s.text(800, 170, 420, 50, '几百万个单元，', 30, IVORY, serif=True)
s.text(800, 215, 420, 50, '由导线连成一张网。', 30, IVORY, serif=True)
bullets(s, 800, 310, 420, ['块内：短导线', '块间：长导线，更慢、更耗电，\n也占布线资源', '目标和通讯录一样：\n块内紧，块间少'], 22)

# 05 ruler and brute force
s = slide('尺子与穷举都不行', '短片 01:25。按位置拿尺子切成三块（左中右各 8 个），会切断权重 88，全部才 202，将近一半。穷举也不行：24 个单元平分成三组已有 24!/(8!³·3!) = 1,577,585,295 种；一百万个单元分成三组约有 3^1,000,000/6 种，这个数字本身就有 477,121 位。两个数字都由项目脚本计算。', '短片 01:25')
big_graph(s, CHIP, fit(CHIP['positions'], (70, 150, 730, 640)), ruler_chip['labels'], cut=True, r=10, weight_scale=.6)
s.text(800, 160, 420, 34, '按位置拿尺子切三块', 22, MUTED)
s.text(800, 196, 420, 70, '切开 88 / 202', 46, GOLD, bold=True, serif=True)
s.text(800, 270, 420, 30, '将近一半的连接被切断（金色）', 18, MUTED)
s.line(800, 330, 1216, 330, LINE, 1)
s.text(800, 352, 420, 34, '试遍所有分法', 22, MUTED)
s.text(800, 392, 420, 50, '24 个单元：15.8 亿种', 30, IVORY, serif=True)
s.text(800, 440, 420, 30, '24!/(8!³·3!) = 1,577,585,295', 17, MUTED)
s.text(800, 494, 420, 50, '一百万个单元：', 30, IVORY, serif=True)
s.text(800, 540, 420, 30, '分法个数约 3^1,000,000/6，本身就有 477,121 位', 17, MUTED)

# 06 one local rule, six questions
s = slide('一条局部规则，六个问题', '短片 01:59。把通讯录的规则借给芯片：每个单元看看邻居们现在的标签，加入连接最强的那一组。规则只有一句话，但留下六个问题。下面九页（第 7–15 页）对应短片下篇的九层，逐个回答。可以先让同学猜每个问题的答案。', '短片 01:59')
s.rect(64, 150, 1152, 130, PANEL, LINE, radius=14)
s.text(100, 172, 1080, 90, '每个点看看邻居们现在的标签，\n加入连接最强的那一组。', 32, GOLD, serif=True, lh=1.45)
qs = [('“最强”怎么算？', '第 2 层'), ('谁先更新？', '第 4、6 层'), ('什么时候停？', '第 5 层'),
      ('它一定会停吗？', '第 6 层'), ('平票怎么办？', '第 2 层'), ('分得好不好？', '第 7、8 层')]
for i, (q, where) in enumerate(qs):
    x, y = 64 + (i % 3) * 388, 320 + (i // 3) * 150
    s.rect(x, y, 364, 126, None, LINE, radius=12)
    s.text(x + 24, y + 22, 320, 44, q, 26, IVORY)
    s.text(x + 24, y + 76, 320, 30, '→ ' + where, 18, TEAL)

# 07 layer 1
init = {n: n for n in 'ABCDE'}
s = slide('第 1 层 · 标签', '短片下篇从这里开始（03:09）。老师的小图：节点名始终是 A–E。开始时每个点的标签就是自己的名字。名字是身份，永远不变；标签是归属，可以改写。后面会看到节点 A 携带标签 B，这不代表 A 改了名。边权是模型中连接的强弱，不是人数，也不自动对应现实的因果强度。孤立点没有邻居，永远保留自己的标签。图和权重来自教师指定的 tiny_graph，边权总和 18。', '短片 03:09')
teacher_graph(s, init)
s.text(800, 170, 420, 50, '名字是身份，', 32, IVORY, serif=True)
s.text(800, 220, 420, 50, '标签是归属。', 32, GOLD, serif=True)
bullets(s, 800, 320, 420, ['开始时，标签 = 自己的名字', 'A 带上标签 B，不等于 A 改名', '边权表示连接强弱，不是人数', '孤立点没有邻居，标签不变'], 21)

# 08 layer 2
wv = FILM['weightVote']
wpos = {'A': (380, 400), 'B': (150, 240), 'C': (150, 560), 'D': (640, 400)}
s = slide('第 2 层 · 怎样算“最强”', 'A 的邻居里，B 和 C 带标签 B（边权各 1），D 带标签 D（边权 4）。数人头，B 有两票；加权重，B 只有 1+1 = 2，D 有 4，所以 A 选 D。人多不一定赢。平票必须有明确规则：课堂默认取字典序最小的候选标签；没有邻居就不改。最大标签、保留当前优先也是可定义的平票方式；随机方式需要记录随机种子。', '短片 03:20')
teacher_graph(s, wv['labels'], edges=[tuple(e) for e in wv['graph']['edges']], pos=wpos, dy=-20)
s.text(800, 160, 420, 40, '按标签把边权加起来，取最大：', 21, MUTED)
s.rect(800, 205, 416, 70, PANEL, LINE, radius=10)
s.text(800, 220, 416, 40, '新标签 = argmax  Σ w(u, v)', 26, IVORY, align='c')
s.table(800, 300, [208, 208], 46, [['标签', '加权票数'], ['B', '1 + 1 = 2'], ['D', '4']], accent={(2, 1): GOLD})
s.text(800, 450, 420, 40, 'A 选 D：人多不一定赢。', 26, GOLD, serif=True)
s.text(800, 510, 420, 90, '平票时取字典序最小的标签；\n没有邻居就不改。', 20, MUTED)

# 09 layer 3
s = slide('第 3 层 · 一次更新', '按 A B C D E 异步访问，轮到 C 时 A、B 已经改为 B。C 读到 B 标签 1 + 4 = 5，D 标签 2，E 标签 1，于是改写为 B，名字仍是 C。一次更新只读这个点自己的边，不需要看到全局。不能讲成“选最粗的那条边”：B 赢是因为两条边的权重加在一起。', '短片 03:40')
teacher_graph(s, c_event['before'], highlight=[('A', 'C'), ('B', 'C')])
s.text(800, 160, 420, 40, '轮到 C，它只读自己的边：', 21, MUTED)
s.table(800, 210, [140, 276], 50, [['标签', '加权票数'], ['B', '1 + 4 = 5'], ['D', '2'], ['E', '1']], accent={(1, 1): GOLD})
s.text(800, 430, 420, 50, 'C 改写为标签 B，', 28, GOLD, serif=True)
s.text(800, 478, 420, 50, '名字仍是 C。', 28, IVORY, serif=True)
s.text(800, 550, 420, 60, '不需要看到全局，\n也不是“选最粗的一条边”。', 19, MUTED)

# 10 layer 4
rows = [['访问', '当时看到的加权票数', '新标签', '改变']]
for e in first_round:
    votes = {}
    for nb in e['neighbors']: votes.setdefault(nb['label'], []).append(nb['weight'])
    seen = '，'.join(f"{L}:{'+'.join(map(str, ws))}{'=' + str(sum(ws)) if len(ws) > 1 else ''}" for L, ws in sorted(votes.items()))
    rows.append([e['node'], seen, e['label'], '是' if e['changed'] else '否'])
s = slide('第 4 层 · 一整轮', '逐行讲第 1 轮。B 立刻读到 A 的新标签：每次更新立即写回，后面的节点读到前面刚产生的新值，这叫异步更新，是课堂默认。第 1 轮结束为 B B B E E，A、C、D 共 3 个节点改变。已经能看出两组，但还没有确认稳定。', '短片 04:00')
s.table(64, 150, [150, 560, 220, 222], 62, rows, size=22, accent={(r, 3): GOLD for r in range(1, 6) if rows[r][3] == '是'})
s.text(64, 548, 700, 50, '轮末标签：B  B  B  E  E', 30, GOLD, serif=True)
s.text(64, 604, 1152, 40, '每次更新立即写回，后面的节点读到新值：这叫异步更新（课堂默认）。', 20, MUTED)

# 11 layer 5
s = slide('第 5 层 · 何时停止', '第 2 轮重新检查 5 个节点，全部不变，才确认稳定，所以轮数是 2（包含确认轮）。只看“这一步没变”不够：后面的点一改，前面的点可能又要改。实现里有三种停止原因：stable（完整一轮没有改变）、cycle（完整状态再次出现）、limit（达到默认 50 轮的安全上限）。不能把 limit 写成已经收敛。', '短片 04:20')
s.table(64, 150, [180, 300, 300], 56, [['轮次', '轮末标签 A B C D E', '改变的节点'], ['第 1 轮', 'B  B  B  E  E', 'A、C、D（3 个）'], ['第 2 轮', 'B  B  B  E  E', '无（0 个）']], size=22, accent={(2, 2): TEAL})
s.text(64, 360, 780, 50, '一整轮零变化，才算停下：共 2 轮。', 28, TEAL, serif=True)
for i, (k, d) in enumerate([('stable', '完整一轮没有任何改变'), ('cycle', '完整的标签状态再次出现'), ('limit', '达到安全上限（默认 50 轮），\n不能写成已经收敛')]):
    x = 64 + i * 388
    s.rect(x, 440, 364, 170, PANEL, LINE, radius=12)
    s.text(x + 24, 462, 320, 44, k, 30, GOLD, bold=True)
    s.text(x + 24, 516, 320, 80, d, 20, IVORY)

# 12 layer 6
s = slide('第 6 层 · 同时还是轮流', '只看一条 A—B 边。同步：本轮全部节点只读本轮开始时的旧标签，轮末统一提交，于是 A/B 变 B/A，再变回 A/B，形成 2 周期振荡，永远停不下来；网页遇到完整状态重复会报告 cycle 并停止。异步：A 先更新为 B，B 读到 A 的新标签，两者一致，下一轮零变化。读取哪个时刻的邻居标签，会改变整个过程。', '短片 04:35')
for i, (x, lab) in enumerate([(294, '同步'), (755, '异步')]):
    s.text(x + 16, 160, 445, 40, lab + ('：大家同时读旧标签' if i == 0 else '：轮流读，立即写回'), 24, GOLD if i == 0 else TEAL, serif=True)
s.table(64, 220, [230, 461, 461], 60, [['', '同步  A / B', '异步  A / B'], ['初始', 'A / B', 'A / B'], ['第 1 轮', 'B / A', 'B / B'], ['第 2 轮', 'A / B（回到初始）', 'B / B（零变化）'], ['停止原因', 'cycle：永远停不下来', 'stable']], size=22,
        accent={(4, 1): GOLD, (4, 2): TEAL})
s.text(64, 560, 1152, 40, '读取哪个时刻的邻居标签，会改变整个过程。', 24, IVORY, serif=True)

# 13 layer 7
s = slide('第 7 层 · 分得好不好', '最终分组 {A,B,C} 与 {D,E}。跨组的是 CD 和 CE 两条边，加权 cutsize = 2 + 1 = 3：加的是权重，不是条数。原图总边权 18，分组规模 3:2，这里没有施加均衡约束。把 C–D 加强到 5，五个点合成一组；加强到 6 或更大（短片演示的是 10），C 被拉到 D、E 那一组，cutsize 变成 5。全放进一组时 cutsize = 0，但未必满足任务：cutsize 小不等于分得好。', '短片 04:55')
teacher_graph(s, GOLDEN['labels'], cut=True)
s.text(800, 160, 420, 40, '跨组的边：CD 与 CE', 22, MUTED)
s.text(800, 200, 420, 60, 'cutsize = 2 + 1 = 3', 38, GOLD, bold=True, serif=True)
s.text(800, 262, 420, 30, '加的是权重，不是条数', 18, MUTED)
s.line(800, 316, 1216, 316, LINE, 1)
s.text(800, 336, 420, 34, '把 C–D 加强：', 22, IVORY)
s.table(800, 380, [150, 266], 46, [['C–D 权重', '分组'], ['2（原图）', '{A,B,C} {D,E}'], ['5', '五个点合成一组'], ['≥ 6（如 10）', '{A,B} {C,D,E}，cut 5']], size=18)
s.text(800, 580, 420, 60, '全放一组 cutsize = 0，\n但未必满足任务。', 19, MUTED)

# 14 layer 8
s = slide('第 8 层 · 回到芯片', '回到 24 个单元，每个从自己的标签出发。颜色沿最强的连线扩散，第 5 轮零变化：三组各 8 个，只切开 2，尺子切开 88。每一轮通过邻接表累计票数，每条连线被两端各读一次，所以一轮的工作量和连线数成正比；整体耗时还取决于轮数、平票处理和输入。不要把论文“近线性时间”理解为任意图都固定几轮。', '短片 05:10')
big_graph(s, CHIP, fit(CHIP['positions'], (70, 150, 730, 640)), lp_chip['labels'], cut=True, r=10, weight_scale=.6)
stat(s, 800, 160, 200, '5 轮', '第 5 轮零变化')
stat(s, 1010, 160, 210, '8 · 8 · 8', '三组大小')
s.text(800, 300, 420, 30, '切开的权重', 20, MUTED)
s.text(800, 336, 200, 60, '2', 52, TEAL, bold=True, serif=True)
s.text(1010, 336, 210, 60, '88', 52, GOLD, bold=True, serif=True)
s.text(800, 400, 200, 30, '标签传播', 17, TEAL)
s.text(1010, 400, 210, 30, '按位置的尺子', 17, GOLD)
s.text(800, 470, 420, 120, '每条连线一轮被两端各读一次：\n一轮的工作量和连线数成正比。', 20, IVORY, lh=1.5)

# 15 layer 9
rev = FILM['reverse']
s = slide('第 9 层 · 它做不到的事', '标签传播不保证分成几组（24 单元图恰好 3 组，老师的图是 2 组，没有参数能指定“分 3 组”），不保证每组一样大（老师的图是 3 + 2），受访问顺序影响，也不是全局最优：基本 LP 不是受约束的最小割求解器。下表是实际独立运行：倒序 E D C B A 得到 A A A D D，含确认轮共 3 轮；分组相同，标签名与过程不同。其他图、权重或规则也可能改变分组本身。', '短片 05:25')
bullets(s, 64, 150, 560, ['不保证分成几组：没有“分 3 组”的参数', '不保证每组一样大：老师的图是 3 + 2', '受访问顺序影响', '不是全局最优：不是受约束的最小割求解器'], 22)
s.table(660, 160, [190, 220, 146], 56, [['访问顺序', '最终标签', '轮数'],
        ['A B C D E', ' '.join(GOLDEN['labels'][n] for n in 'ABCDE'), str(GOLDEN['rounds'])],
        ['E D C B A', ' '.join(rev['labels'][n] for n in 'ABCDE'), str(rev['rounds'])]], size=20)
s.text(660, 350, 556, 90, '分组一样（三点 + 两点），\n标签名字和过程不同。', 22, TEAL, serif=True, lh=1.45)
s.text(660, 460, 556, 80, '比较结果时，要分清“标签叫什么”\n和“哪些点在一起”。', 19, MUTED)

# 16 real use
s = slide('现实中的用法', '短片 05:43。在芯片设计中，聚类常用来先把紧密的单元收成团，让电路图变小（粗化），再交给更精细的划分、布局与布线；那里还要处理面积、时序、拥塞等约束。真实网表也可能采用超图模型。本实验不替代商业 EDA 流程。生活中的类比只解释图结构，不能据此判断个人本质或推出因果。结尾一句：没有人画圈，圈子从连接里长出来。', '短片 05:43')
steps = [('网表', '几百万单元'), ('标签传播', '收成团（粗化）'), ('精细划分', '面积 · 平衡'), ('布局布线', '时序 · 拥塞')]
for i, (a, b) in enumerate(steps):
    x = 64 + i * 292
    s.rect(x, 170, 252, 130, PANEL, TEAL if i == 1 else LINE, 2 if i == 1 else 1.5, radius=12)
    s.text(x, 196, 252, 44, a, 28, TEAL if i == 1 else IVORY, align='c', serif=True)
    s.text(x, 248, 252, 30, b, 18, MUTED, align='c')
    if i < 3: s.text(x + 252, 214, 40, 40, '→', 26, FAINT, align='c')
s.text(64, 336, 1152, 40, '它给出线索，不是最终答案；真实网表还可能用超图建模。', 20, MUTED)
s.text(64, 440, 1152, 70, '没有人画圈。', 48, IVORY, serif=True, align='c')
s.text(64, 520, 1152, 70, '圈子，从连接里长出来。', 48, GOLD, serif=True, align='c')

# 17 classroom
s = slide('课堂互动', '打开网页（地址见本页）。建议顺序：先在“九层讲解”里点按钮跳到短片对应时间；再进“闯关”，九关对应九层，外加两场 Boss；最后在自由实验台修改 C–D 权重、切换同步模式、比较访问顺序。不需要账号，算法在浏览器里运行，进度只存在本机。', '网页')
acts = [('1', '预测 C 的下一个标签', '说出 B、D、E 各得几票，再单步验证（第 3 关）'),
        ('2', '让两节点同步运行', '观察为什么停不下来，再切回异步（第 6 关）'),
        ('3', '把 C–D 调到 5 和 10', '看边界怎样移动、cutsize 怎样变（第 8 关）'),
        ('4', 'Boss：给芯片分块', '亲手分成三组，比尺子的 88 分得更好')]
for i, (n, a, b) in enumerate(acts):
    y = 150 + i * 112
    s.circle(100, y + 40, 26, None, GOLD, 2)
    s.text(74, y + 22, 52, 40, n, 24, GOLD, bold=True, align='c')
    s.text(150, y + 8, 1060, 40, a, 26, IVORY)
    s.text(150, y + 50, 1060, 30, b, 18, MUTED)
s.text(64, 612, 1152, 34, SITE, 20, TEAL)

# 18 sources
s = slide('资料与许可', '来源完整链接见本页。教师文件没有独立开源许可声明，因此本项目不分发其课件、视频或课程私密资料，只引用事实与小图数据并注明来源。原创代码采用 MIT，原创教材、图解与视频画面采用 CC BY 4.0；第三方音乐、工具和字体保留原许可，见 THIRD_PARTY_NOTICES.md。', '')
s.text(64, 150, 1152, 34, '课程 Clustering Lab05：README、EXAMPLES、CHECKLIST、参考求解器与 tiny_graph', 22, IVORY)
s.text(64, 188, 1152, 30, TEACHER, 15, MUTED)
s.text(64, 250, 1152, 34, 'Raghavan, Albert, Kumara (2007). Near linear time algorithm to detect community', 22, IVORY)
s.text(64, 284, 1152, 34, 'structures in large-scale networks. Physical Review E 76, 036106.', 22, IVORY)
s.text(64, 322, 1152, 30, PAPER, 15, MUTED)
s.text(64, 384, 1152, 34, '配乐：《We Never (Remix)（若如初见）》Hi Noise / 秋成，只随成片分发，不属于本项目许可。', 20, IVORY)
s.line(64, 456, 1216, 456, LINE, 1)
s.text(64, 480, 1152, 34, '原创代码 MIT · 原创教材、图解与视频画面 CC BY 4.0', 22, GOLD)
s.text(64, 524, 1152, 34, '所有标签、票数、轮数与 cutsize 都由项目算法实际运行得出。', 20, MUTED)

for i, s in enumerate(SLIDES, 1): frame(s, i)


# ---------------------------------------------------------------- cairo backend (PNG + PDF)
def wrap(ctx, s, size, w):
    out = []
    for para in s.split('\n'):
        line = ''
        for ch in para:
            if line and ctx.text_extents(line + ch).x_advance > w: out.append(line); line = ch
            else: line += ch
        out.append(line)
    return out


def draw_cairo(ctx, sl):
    ctx.set_source_rgb(*rgb(BG)); ctx.paint()
    for kind, k in sl.items:
        if kind == 'image':
            img = cairo.ImageSurface.create_from_png(str(png_of(k['path'])))
            ctx.save(); ctx.translate(k['x'], k['y']); ctx.scale(k['w'] / img.get_width(), k['h'] / img.get_height())
            ctx.set_source_surface(img, 0, 0); ctx.paint(); ctx.restore()
        elif kind == 'line':
            ctx.set_source_rgb(*rgb(k['color'])); ctx.set_line_width(k['width']); ctx.set_line_cap(cairo.LINE_CAP_ROUND)
            ctx.move_to(k['x1'], k['y1']); ctx.line_to(k['x2'], k['y2']); ctx.stroke()
        elif kind == 'circle':
            ctx.new_path(); ctx.arc(k['cx'], k['cy'], k['r'], 0, 2 * math.pi)
            if k['fill']: ctx.set_source_rgb(*rgb(k['fill'])); ctx.fill_preserve()
            if k['stroke']: ctx.set_source_rgb(*rgb(k['stroke'])); ctx.set_line_width(k['width']); ctx.stroke()
            ctx.new_path()
        elif kind == 'rect':
            rounded(ctx, k['x'], k['y'], k['w'], k['h'], k['radius'])
            if k['fill']: ctx.set_source_rgb(*rgb(k['fill'])); ctx.fill_preserve()
            if k['stroke']: ctx.set_source_rgb(*rgb(k['stroke'])); ctx.set_line_width(k['width']); ctx.stroke()
            ctx.new_path()
        elif kind == 'text':
            cairo_text(ctx, k)
        elif kind == 'table':
            x0, y = k['x'], k['y']
            for r, row in enumerate(k['rows']):
                x = x0
                for c, cell in enumerate(row):
                    w = k['widths'][c]
                    ctx.rectangle(x, y, w, k['rowh'])
                    ctx.set_source_rgb(*rgb(mix(GOLD, .16) if r == 0 else (PANEL if r % 2 else BG))); ctx.fill_preserve()
                    ctx.set_source_rgb(*rgb(LINE)); ctx.set_line_width(1); ctx.stroke()
                    col = GOLD if r == 0 else k['accent'].get((r, c), IVORY)
                    cairo_text(ctx, dict(x=x + 16, y=y, w=w - 24, h=k['rowh'], s=str(cell), size=k['size'], color=col,
                                         bold=r == 0, serif=False, align='l', valign='m', lh=1.2))
                    x += w
                y += k['rowh']


def rounded(ctx, x, y, w, h, r):
    ctx.new_path(); ctx.arc(x + w - r, y + r, r, -math.pi / 2, 0); ctx.arc(x + w - r, y + h - r, r, 0, math.pi / 2)
    ctx.arc(x + r, y + h - r, r, math.pi / 2, math.pi); ctx.arc(x + r, y + r, r, math.pi, 3 * math.pi / 2); ctx.close_path()


def cairo_text(ctx, k):
    ctx.select_font_face(SERIF_PNG if k['serif'] else SANS_PNG, cairo.FONT_SLANT_NORMAL,
                         cairo.FONT_WEIGHT_BOLD if k['bold'] else cairo.FONT_WEIGHT_NORMAL)
    ctx.set_font_size(k['size']); ctx.set_source_rgb(*rgb(k['color']))
    lines = wrap(ctx, k['s'], k['size'], k['w'])
    step = k['size'] * k['lh']; total = step * len(lines)
    y = k['y'] + (k['h'] - total) / 2 if k['valign'] == 'm' else k['y']
    for ln in lines:
        adv = ctx.text_extents(ln).x_advance
        x = k['x'] + (k['w'] - adv if k['align'] == 'r' else (k['w'] - adv) / 2 if k['align'] == 'c' else 0)
        ctx.move_to(x, y + (step - k['size']) / 2 + k['size'] * .88); ctx.show_text(ln); y += step


_png_cache = {}
def png_of(path):
    if path not in _png_cache:
        out = ROOT / 'build/private/deck' / (Path(path).stem + '.png'); out.parent.mkdir(parents=True, exist_ok=True)
        subprocess.run(['ffmpeg', '-v', 'error', '-y', '-i', str(ROOT / path), str(out)], check=True)
        _png_cache[path] = out
    return _png_cache[path]


# ---------------------------------------------------------------- python-pptx backend
def E(v): return Emu(round(v * EMU))
def color(c): return RGBColor.from_string(c.lstrip('#'))


def set_font(run, k):
    f = run.font; f.size = Pt(k['size'] * .75); f.bold = k['bold']; f.color.rgb = color(k['color'])
    face = SERIF_PPT if k['serif'] else SANS_PPT; f.name = face
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None: el = rpr.makeelement(qn(tag), {}); rpr.append(el)
        el.set('typeface', face)


def fill_frame(tf, k):
    tf.word_wrap = True; tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if k['valign'] == 'm' else MSO_ANCHOR.TOP
    for i, para in enumerate(k['s'].split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[k['align']]
        p.line_spacing = k['lh'] * .85
        set_font(p.add_run(), k); p.runs[0].text = para


def draw_pptx(prs, sl):
    page = prs.slides.add_slide(prs.slide_layouts[6])
    page.background.fill.solid(); page.background.fill.fore_color.rgb = color(BG)
    sh = page.shapes
    for kind, k in sl.items:
        if kind == 'image':
            sh.add_picture(str(ROOT / k['path']), E(k['x']), E(k['y']), E(k['w']), E(k['h']))
        elif kind == 'line':
            ln = sh.add_connector(MSO_CONNECTOR.STRAIGHT, E(k['x1']), E(k['y1']), E(k['x2']), E(k['y2']))
            ln.line.color.rgb = color(k['color']); ln.line.width = E(k['width'])
        elif kind in ('circle', 'rect'):
            if kind == 'circle':
                shp = sh.add_shape(MSO_SHAPE.OVAL, E(k['cx'] - k['r']), E(k['cy'] - k['r']), E(2 * k['r']), E(2 * k['r']))
            else:
                shp = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(k['x']), E(k['y']), E(k['w']), E(k['h']))
                shp.adjustments[0] = min(.5, k['radius'] / min(k['w'], k['h']))
            shp.shadow.inherit = False
            if k['fill']: shp.fill.solid(); shp.fill.fore_color.rgb = color(k['fill'])
            else: shp.fill.background()
            if k['stroke']: shp.line.color.rgb = color(k['stroke']); shp.line.width = E(k['width'])
            else: shp.line.fill.background()
        elif kind == 'text':
            tb = sh.add_textbox(E(k['x']), E(k['y']), E(k['w']), E(k['h'])); fill_frame(tb.text_frame, k)
        elif kind == 'table':
            rows, widths = k['rows'], k['widths']
            gt = sh.add_table(len(rows), len(widths), E(k['x']), E(k['y']), E(sum(widths)), E(k['rowh'] * len(rows)))
            tbl = gt.table
            tbl.first_row = True; tbl.horz_banding = False
            for c, w in enumerate(widths): tbl.columns[c].width = E(w)
            for r in range(len(rows)):
                tbl.rows[r].height = E(k['rowh'])
                for c in range(len(widths)):
                    cell = tbl.cell(r, c); cell.fill.solid()
                    cell.fill.fore_color.rgb = color(mix(GOLD, .16) if r == 0 else (PANEL if r % 2 else BG))
                    cell.margin_left = E(16); cell.margin_right = E(8); cell.margin_top = cell.margin_bottom = 0
                    tf = cell.text_frame; tf.text = ''
                    fill_frame(tf, dict(s=str(rows[r][c]), size=k['size'], color=GOLD if r == 0 else k['accent'].get((r, c), IVORY),
                                        bold=r == 0, serif=False, align='l', valign='m', lh=1.2))
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    page.notes_slide.notes_text_frame.text = sl.notes


# ---------------------------------------------------------------- build
def main():
    notes = [{'slide': i, 'title': s.title or '圈子从哪里来', 'notes': s.notes + f'\n\n资料：{TEACHER}\n{PAPER}'} for i, s in enumerate(SLIDES, 1)]
    for s, n in zip(SLIDES, notes): s.notes = n['notes']
    out_png = ROOT / 'site/slides'; out_png.mkdir(parents=True, exist_ok=True)
    for old in out_png.glob('slide-*.png'): old.unlink()
    pdf = cairo.PDFSurface(str(ROOT / 'site/downloads/label-propagation.pdf'), W, H)
    pdf.set_metadata(cairo.PDFMetadata.TITLE, '圈子从哪里来 · 标签传播课堂课件'); pdf.set_metadata(cairo.PDFMetadata.AUTHOR, 'Su3037')
    pctx = cairo.Context(pdf)
    for i, sl in enumerate(SLIDES, 1):
        draw_cairo(pctx, sl); pctx.show_page()
        img = cairo.ImageSurface(cairo.FORMAT_RGB24, round(W * PNG_SCALE), round(H * PNG_SCALE))
        ctx = cairo.Context(img); ctx.scale(PNG_SCALE, PNG_SCALE); draw_cairo(ctx, sl)
        tmp = ROOT / 'build/private/deck' / f'slide-{i:02d}.png'; tmp.parent.mkdir(parents=True, exist_ok=True)
        img.write_to_png(str(tmp))
        subprocess.run(['ffmpeg', '-v', 'error', '-y', '-i', str(tmp), '-vf', 'scale=1600:-1', '-pix_fmt', 'rgb24',
                        str(out_png / f'slide-{i:02d}.png')], check=True)
    pdf.finish()
    prs = Presentation(); prs.slide_width, prs.slide_height = E(W), E(H)
    prs.core_properties.title = '圈子从哪里来 · 标签传播课堂课件'; prs.core_properties.author = 'Su3037'
    for sl in SLIDES: draw_pptx(prs, sl)
    prs.save(str(ROOT / 'site/downloads/label-propagation.pptx'))
    text = json.dumps(notes, ensure_ascii=False, indent=2) + '\n'
    (ROOT / 'presentation').mkdir(exist_ok=True)
    for p in (ROOT / 'presentation/notes.json', ROOT / 'site/downloads/notes.json'): p.write_text(text, encoding='utf-8')
    print(f'Built {len(SLIDES)} slides: PPTX, PDF, PNG previews and notes')


if __name__ == '__main__':
    main()
